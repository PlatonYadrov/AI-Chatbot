from __future__ import annotations

"""Convert PowerPoint presentations into slide level blocks."""

import logging
from typing import Iterator, Optional

from .base import BaseParser, RawBlock, ensure_dependency
from .ocr import OcrEngine, merge_ocr_text

LOGGER = logging.getLogger(__name__)


class PptxParser(BaseParser):
    """Convert PowerPoint presentations into slide level blocks."""

    def __init__(self, *, ocr_engine: Optional[OcrEngine] = None) -> None:
        self.ocr_engine = ocr_engine or OcrEngine()

    def parse(self, path: str, *, doc_id: Optional[str] = None) -> Iterator[RawBlock]:
        ensure_dependency("pptx", "pip install python-pptx")
        from pptx import Presentation  # type: ignore

        resolved_id = self._resolve_doc_id(path, doc_id)
        LOGGER.debug("Parsing PPTX document", extra={"doc_id": resolved_id, "path": path})

        try:
            presentation = Presentation(path)
        except Exception:
            LOGGER.exception("pptx_open_failed", extra={"path": path})
            raise
        for slide_index, slide in enumerate(presentation.slides, start=1):
            fragments: list[str] = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "") or ""
                text = self._normalise(text)
                if text:
                    fragments.append(text)

            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = self._normalise(slide.notes_slide.notes_text_frame.text or "")
                if notes_text:
                    fragments.append(f"Notes: {notes_text}")

            payload = "\n\n".join(fragments)
            payload = self._attach_ocr(slide, payload)
            if not payload:
                continue

            meta = {
                "doc_id": resolved_id,
                "type": "pptx",
                "path": path,
                "slide": slide_index,
            }
            if payload != "\n\n".join(fragments):
                meta["has_ocr"] = True
                meta["image_count"] = meta.get("image_count", 0) + 1

            yield RawBlock(text=payload, meta=meta)

    def _attach_ocr(self, slide, payload: str) -> str:
        if payload is None:
            payload = ""
        if not self.ocr_engine.is_available():
            return payload

        ocr_results: list[str] = []
        for shape in slide.shapes:
            image = getattr(shape, "image", None)
            blob = getattr(image, "blob", None)
            if not blob:
                continue
            text = self.ocr_engine.extract_text(blob)
            if text:
                ocr_results.append(text)

        if not ocr_results:
            return payload
        return merge_ocr_text(payload, ocr_text="\n\n".join(ocr_results))


__all__ = ["PptxParser"]